import io
from pathlib import Path
from uuid import uuid4

from fastapi import APIRouter, File, Form, HTTPException, UploadFile, status

from app.config import ALLOWED_EXTENSIONS, MAX_FILE_SIZE, UPLOAD_DIR
from app.deps import engine
from app.schemas import UploadResponseSchema

router = APIRouter(prefix="/api", tags=["Документы"])


def extract_text(ext: str, content: bytes) -> str:
    if ext == ".docx":
        import docx
        d = docx.Document(io.BytesIO(content))
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            parts += [" | ".join(c.text for c in r.cells) for r in t.rows]
        return "\n".join(parts)
    if ext == ".pdf":
        import pypdf
        return "\n".join((p.extract_text() or "") for p in pypdf.PdfReader(io.BytesIO(content)).pages)
    return content.decode("utf-8", errors="ignore")


@router.post("/upload", response_model=UploadResponseSchema, status_code=status.HTTP_201_CREATED, summary="Загрузить проект документа для проверки")
async def upload_document(file: UploadFile = File(...), session_id: str | None = Form(default=None)) -> UploadResponseSchema:
    filename = file.filename or ""
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(400, "Можно загрузить только PDF, DOCX, TXT или MD")
    content = await file.read()
    if not content:
        raise HTTPException(400, "Файл пустой")
    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(400, "Максимальный размер — 10 МБ")
    try:
        text = extract_text(ext, content)
    except Exception:
        raise HTTPException(400, "Не удалось прочитать файл")
    if not text.strip():
        raise HTTPException(400, "В файле не найден текст (скан без распознавания?)")
    doc_id = str(uuid4())
    UPLOAD_DIR.mkdir(exist_ok=True)
    (UPLOAD_DIR / f"{doc_id}{ext}").write_bytes(content)
    sid = engine.attach_document(session_id, filename, text)
    return UploadResponseSchema(document_id=doc_id, filename=filename, status="uploaded", session_id=sid, chars=len(text))
