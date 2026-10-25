"""База знаний: загрузка документов, нарезка на положения (chunks), версионность, доступ."""
from __future__ import annotations

import re
import threading
from dataclasses import dataclass, field
from pathlib import Path

from app.config import KB_DIR

HEAD_RE = re.compile(r"^##\s*\[(?P<id>[^\]]+)\]\s*(?P<title>[^|\n]*)(?:\|\s*(?P<ref>.*))?$")
NUM_RE = re.compile(r"^\s*(\d+(?:\.\d+){0,3})[.)]?\s+(?=\S)")


@dataclass
class Chunk:
    id: str
    title: str
    ref: str          # человекочитаемая ссылка: «Матрица визирования, п. 3»
    text: str
    document: str
    version: str = "1"
    status: str = "active"      # active | archived
    access: str = "public"      # public | restricted
    stems: list[str] = field(default_factory=list)


def _front_matter(raw: str) -> tuple[dict, str]:
    meta: dict = {}
    if raw.startswith("---"):
        parts = raw.split("---", 2)
        if len(parts) >= 3:
            for line in parts[1].splitlines():
                if ":" in line:
                    k, v = line.split(":", 1)
                    meta[k.strip()] = v.strip()
            raw = parts[2]
    return meta, raw


def _parse_md(raw: str, fallback_doc: str) -> list[Chunk]:
    meta, body = _front_matter(raw)
    body = re.sub(r"<!--.*?-->", "", body, flags=re.S)
    doc = meta.get("doc", fallback_doc)
    chunks, cur = [], None
    for line in body.splitlines():
        m = HEAD_RE.match(line.strip())
        if m:
            if cur:
                chunks.append(cur)
            cur = Chunk(id=m["id"].strip(), title=m["title"].strip(), ref=(m["ref"] or doc).strip(), text="",
                        document=doc, version=meta.get("version", "1"), status=meta.get("status", "active"),
                        access=meta.get("access", "public"))
        elif cur is not None:
            cur.text += line + "\n"
    if cur:
        chunks.append(cur)
    for c in chunks:
        c.text = c.text.strip()
    return chunks


def _read_plain(path: Path) -> str:
    ext = path.suffix.lower()
    if ext == ".docx":
        import docx
        doc = docx.Document(str(path))
        rows = [" | ".join(c.text for c in row.cells) for table in doc.tables for row in table.rows]
        return "\n".join([p.text for p in doc.paragraphs] + rows)
    if ext == ".pdf":
        import pypdf
        return "\n".join((p.extract_text() or "") for p in pypdf.PdfReader(str(path)).pages)
    return path.read_text(encoding="utf-8", errors="ignore")


def _parse_plain(path: Path) -> list[Chunk]:
    """docx/pdf/txt: режем по нумерованным пунктам (1., 1.2, 4.3 ...). Нет нумерации — режем по абзацам."""
    text = _read_plain(path)
    doc, slug = path.stem, re.sub(r"\W+", "_", path.stem.lower()).strip("_")[:24]
    chunks: list[Chunk] = []
    cur_num, buf = None, []

    def flush():
        if cur_num and buf:
            t = " ".join(buf).strip()
            chunks.append(Chunk(id=f"{slug}_{cur_num}", title=t[:70], ref=f"{doc}, п. {cur_num}", text=t, document=doc))

    for line in text.splitlines():
        m = NUM_RE.match(line)
        if m:
            flush()
            cur_num, buf = m.group(1), [line[m.end():].strip()]
        elif line.strip():
            buf.append(line.strip())
    flush()
    if not chunks:
        for i, para in enumerate(p for p in re.split(r"\n\s*\n", text) if p.strip()):
            chunks.append(Chunk(id=f"{slug}_p{i+1}", title=para[:70], ref=f"{doc}, абзац {i+1}", text=para.strip(), document=doc))
    return chunks


class KnowledgeBase:
    def __init__(self, kb_dir: Path = KB_DIR):
        self.kb_dir = kb_dir
        self._lock = threading.Lock()
        self.chunks: list[Chunk] = []
        self.reload()

    def reload(self) -> int:
        from app.services.nlp import content_stems
        loaded: list[Chunk] = []
        # rglob позволяет хранить исходные материалы и догруженные файлы по папкам.
        for path in sorted(self.kb_dir.rglob("*")):
            if path.suffix.lower() not in {".md", ".txt", ".docx", ".pdf"} or path.name.startswith("."):
                continue
            try:
                items = _parse_md(path.read_text(encoding="utf-8"), path.stem) if path.suffix.lower() == ".md" else _parse_plain(path)
            except Exception as exc:   # битый файл не должен ронять сервис
                print(f"[kb] пропущен {path.name}: {exc}")
                continue
            loaded += items
        # версионность: у одного и того же chunk id побеждает последняя active-версия
        by_id: dict[str, Chunk] = {}
        for c in loaded:
            if c.status != "active":
                continue
            old = by_id.get(c.id)
            if old is None or c.version >= old.version:
                by_id[c.id] = c
        for c in by_id.values():
            c.stems = content_stems(f"{c.title} {c.title} {c.ref} {c.text}", expand=False)
        with self._lock:
            self.chunks = list(by_id.values())
        return len(self.chunks)

    def get(self, chunk_id: str) -> Chunk | None:
        return next((c for c in self.chunks if c.id == chunk_id), None)

    def info(self) -> dict:
        docs: dict[str, dict] = {}
        for c in self.chunks:
            d = docs.setdefault(c.document, {"document": c.document, "version": c.version, "chunks": 0})
            d["chunks"] += 1
        return {"chunks": len(self.chunks), "documents": list(docs.values())}
